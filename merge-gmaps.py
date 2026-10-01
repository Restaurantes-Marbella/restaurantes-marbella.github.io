# Cruza _data-portal.json (geocodificación por portal) con _gmaps.json (Google Maps)
# → data.json para la web + Excel con columnas de coordenadas.
import json, re, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

base = json.load(open('_data-portal.json', encoding='utf-8'))
gm = json.load(open('_gmaps.json', encoding='utf-8'))

VACIAS = set('la el los las de del y the bar restaurante restaurant cafeteria cafe taberna marbella casco antiguo puente malaga '
             'pizzeria heladeria churreria bodega gastrobar tapas & - | by'.split())
def tokens(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower()
    return {t for t in re.findall(r'[a-z0-9]+', s) if t not in VACIAS and len(t) > 1}
def plano(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)
def coincide(a, b):
    from difflib import SequenceMatcher
    pa, pb = plano(a), plano(b)
    # "T&M" ⊂ "T&M Kebab", "Casablanca" ≈ "Casa Blanca"
    if len(pa) >= 2 and pa in pb: return True
    ta, tb = tokens(a), tokens(b)
    # Erratas: "Primeira" ≈ "Primeria"
    if any(len(x) >= 5 and SequenceMatcher(None, x, y).ratio() >= 0.85 for x in ta for y in tb): return True
    if not ta or not tb: return False
    if ta & tb: return True
    ja, jb = ''.join(sorted(ta)), ''.join(sorted(tb))
    return ja in ''.join(tb) or jb in ''.join(ta)

# Google devuelve el lugar que contiene al local (hotel, Mercado): la ubicación es válida
CONTENEDOR = {17, 278, 284}

FUENTE = {'google': 'Google Maps', 'portal': 'Portal (CartoCiudad/OSM)', 'calle': 'Punto de calle (aprox.)'}
out, filas, cuenta = [], [], {}
for r in base:
    g = gm.get(str(r['id']), {})
    n = dict(r)
    if g.get('estado') == 'ok' and (coincide(r['nombre'], g.get('titulo')) or r['id'] in CONTENEDOR):
        n['lat'], n['lng'], n['precision'] = g['lat'], g['lng'], 'google'
        estado = 'OK'
    elif g.get('estado') == 'ok':
        estado = 'REVISAR: Google devuelve otro nombre'
    elif g.get('estado') == 'fuera_de_zona':
        estado = 'REVISAR: Google lo sitúa fuera del centro'
    else:
        estado = 'No encontrado en Google'
    cuenta[estado] = cuenta.get(estado, 0) + 1
    out.append(n)
    filas.append([r['id'], r['zona'], r['calle'], r['num'], r['nombre'], r['tipo'], r['categoria'],
                  'Sí' if r['accesible'] == 'si' else 'No', r['obs'],
                  n['lat'], n['lng'], FUENTE[n['precision']], g.get('titulo', ''), estado,
                  f"https://www.google.com/maps?q={n['lat']},{n['lng']}"])

# Separar levemente (~4 m) locales que comparten exactamente el mismo punto, solo en la web
seen = {}
for n in out:
    k = (round(n['lat'], 6), round(n['lng'], 6))
    i = seen[k] = seen.get(k, 0) + 1
    if i > 1:
        import math
        a = i * 2.4
        n['lat'] = round(n['lat'] + 0.00004 * math.cos(a), 6)
        n['lng'] = round(n['lng'] + 0.00005 * math.sin(a), 6)
json.dump(out, open('data.json', 'w', encoding='utf-8'), ensure_ascii=False)

wb = Workbook(); ws = wb.active; ws.title = 'Restaurantes'
cab = ['Nº', 'Zona', 'Calle', 'nº', 'Nombre', 'Tipo', 'Categoría', '¿Es accesible?', 'Observaciones',
       'Latitud', 'Longitud', 'Fuente coordenadas', 'Nombre en Google Maps', 'Estado', 'Ver en mapa']
ws.append(cab)
for f in filas: ws.append(f)
azul = PatternFill('solid', fgColor='1E4E8C'); amar = PatternFill('solid', fgColor='FFF2CC')
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = azul; c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    row[9].number_format = row[10].number_format = '0.0000000'
    row[14].hyperlink = row[14].value; row[14].value = 'Abrir'; row[14].font = Font(color='0563C1', underline='single')
    if row[13].value != 'OK':
        for c in row: c.fill = amar
for col, w in zip('ABCDEFGHIJKLMNO', [5, 22, 30, 10, 30, 13, 28, 12, 25, 12, 12, 24, 34, 38, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wb.save('Restaurantes Marbella - coordenadas.xlsx')
print(cuenta)
print('\n'.join(f"{f[0]} {f[4]} → {f[12]} | {f[13]}" for f in filas if f[13] != 'OK'))
